#!/usr/bin/env python3
"""Validates the files that test/core/formats_test.dart writes to
build/formats with independent readers: ezdxf (DXF), LibreDWG (DWG),
lxml (KML, SVG), json (GeoJSON, project, catalogue pack), openpyxl (XLSX),
pypdf (PDF), csv and zipfile.

    flutter test test/core/formats_test.dart
    python3 tool/formats/check_exports.py      # pip install ezdxf openpyxl lxml pypdf

Exit code 1 when any check fails.
"""
import csv
import io
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import ezdxf
import openpyxl
from ezdxf import recover
from lxml import etree
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'build' / 'formats'
failures = []


def check(cond, msg):
    if cond:
        print(f'  ok   {msg}')
    else:
        print(f'  FAIL {msg}')
        failures.append(msg)


def dxf_summary(path):
    doc, auditor = recover.readfile(path)
    types = {}
    for e in doc.modelspace():
        types[e.dxftype()] = types.get(e.dxftype(), 0) + 1
    return doc, auditor, types


def check_dxf(path):
    doc, auditor, types = dxf_summary(path)
    check(not auditor.has_errors, f'{path.name}: ezdxf audit has no errors ({len(auditor.errors)} errors, {len(auditor.fixes)} fixes) {[e.message for e in auditor.errors[:3]]}')
    check(sum(types.values()) > 0, f'{path.name}: {sum(types.values())} entities {types}')
    check(len(doc.layers) > 1, f'{path.name}: {len(doc.layers)} layers')
    check(doc.header.get('$INSUNITS', 0) == 6 or doc.dxfversion == 'AC1009', f'{path.name}: units are metres ($INSUNITS={doc.header.get("$INSUNITS")}, {doc.dxfversion})')
    return types


def check_dwg(path):
    if not shutil.which('dwgread'):
        print('  skip dwgread not installed')
        return
    r = subprocess.run(['dwgread', '-v2', str(path)], capture_output=True, text=True)
    log = r.stdout + r.stderr
    bad = [l for l in log.splitlines() if 'ERROR' in l or 'mismatch' in l]
    check(r.returncode == 0 and not bad, f'{path.name}: dwgread reads it without errors {bad[:3]}')
    out = OUT / (path.stem + '_libredwg.dxf')
    if out.exists():
        out.unlink()
    r = subprocess.run(['dwg2dxf', '-y', '-o', str(out), str(path)], capture_output=True, text=True)
    check(r.returncode == 0 and out.exists(), f'{path.name}: dwg2dxf converts it')
    if out.exists():
        doc, auditor, types = dxf_summary(out)
        check(sum(types.values()) > 0, f'{path.name}: LibreDWG DXF has {sum(types.values())} entities {types}')
        return types


KML_NS = '{http://www.opengis.net/kml/2.2}'


def check_kml_text(name, text):
    root = etree.fromstring(text)
    check(root.tag == f'{KML_NS}kml', f'{name}: root is kml 2.2')
    placemarks = root.findall(f'.//{KML_NS}Placemark')
    check(len(placemarks) > 0, f'{name}: {len(placemarks)} placemarks in {len(root.findall(f".//{KML_NS}Folder"))} folders')
    styles = {s.get('id') for s in root.findall(f'.//{KML_NS}Style')}
    refs = {u.text.lstrip('#') for u in root.findall(f'.//{KML_NS}styleUrl')}
    check(refs <= styles, f'{name}: every styleUrl resolves ({len(styles)} styles)')
    bad = 0
    for coords in root.findall(f'.//{KML_NS}coordinates'):
        for tok in coords.text.split():
            lon, lat = (float(v) for v in tok.split(',')[:2])
            if not (-180 <= lon <= 180 and -90 <= lat <= 90):
                bad += 1
    check(bad == 0, f'{name}: every coordinate is lon,lat within range')
    rings = root.findall(f'.//{KML_NS}LinearRing/{KML_NS}coordinates')
    closed = all(r.text.split()[0] == r.text.split()[-1] for r in rings)
    check(closed, f'{name}: {len(rings)} polygon rings closed')
    return len(placemarks)


def check_geojson(path):
    j = json.loads(path.read_text(encoding='utf-8'))
    check(j.get('type') == 'FeatureCollection', f'{path.name}: FeatureCollection')
    feats = j.get('features', [])
    bad = []
    for f in feats:
        g = f.get('geometry', {})
        t, c = g.get('type'), g.get('coordinates')
        if f.get('type') != 'Feature' or 'properties' not in f:
            bad.append('feature shape')
        if t == 'Point':
            ok = 2 <= len(c) <= 3 and -180 <= c[0] <= 180 and -90 <= c[1] <= 90
        elif t == 'LineString':
            ok = len(c) >= 2 and all(-180 <= p[0] <= 180 and -90 <= p[1] <= 90 for p in c)
        elif t == 'Polygon':
            ok = all(len(r) >= 4 and r[0] == r[-1] for r in c)
        else:
            ok = False
        if not ok:
            bad.append(t)
    check(not bad, f'{path.name}: {len(feats)} features valid per RFC 7946 {bad[:3]}')
    return len(feats)


def check_svg(path):
    root = etree.fromstring(path.read_bytes())
    check(root.tag == '{http://www.w3.org/2000/svg}svg', f'{path.name}: root is svg')
    vb = root.get('viewBox').split()
    check(len(vb) == 4 and float(vb[2]) > 0 and float(vb[3]) > 0, f'{path.name}: viewBox {vb}')
    groups = root.findall('{http://www.w3.org/2000/svg}g')
    paths = root.findall('.//{http://www.w3.org/2000/svg}path')
    texts = root.findall('.//{http://www.w3.org/2000/svg}text')
    check(groups and paths and texts, f'{path.name}: {len(groups)} layer groups, {len(paths)} paths, {len(texts)} texts')


def check_pdf(path, expect_text):
    r = PdfReader(str(path))
    text = ''.join(p.extract_text() or '' for p in r.pages)
    check(len(r.pages) >= 1, f'{path.name}: {len(r.pages)} pages, title {r.metadata.title!r}')
    for t in expect_text:
        check(t in text, f'{path.name}: text contains {t!r}')


def check_xlsx(path, expect_sheets):
    wb = openpyxl.load_workbook(str(path))
    check(len(wb.sheetnames) >= expect_sheets, f'{path.name}: sheets {wb.sheetnames}')
    ws = wb.worksheets[0]
    header_bold = all(c.font.b for c in ws[1] if c.value is not None)
    check(header_bold, f'{path.name}: header row bold')
    nums = sum(1 for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, (int, float)))
    check(nums > 0, f'{path.name}: {nums} numeric cells in the first sheet')


def check_csv(path, delimiter=';', min_rows=2):
    text = path.read_text(encoding='utf-8-sig')
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    check(len(rows) >= min_rows, f'{path.name}: {len(rows)} rows, first {rows[0][:4]}')


def main():
    if not OUT.exists():
        print('build/formats missing: run flutter test test/core/formats_test.dart first')
        sys.exit(1)
    print('drawing exports')
    dxf_types = check_dxf(OUT / 'orchard.dxf')
    dwg_types = check_dwg(OUT / 'orchard.dwg')
    if dwg_types:
        n_dxf = sum(dxf_types.values())
        n_dwg = sum(dwg_types.values())
        check(abs(n_dxf - n_dwg) <= n_dxf * 0.05, f'DXF ({n_dxf}) and DWG ({n_dwg}) carry the same entities')
    n_kml = check_kml_text('orchard.kml', (OUT / 'orchard.kml').read_bytes())
    with zipfile.ZipFile(OUT / 'orchard.kmz') as z:
        check(z.testzip() is None and 'doc.kml' in z.namelist(), 'orchard.kmz: valid zip with doc.kml')
        n_kmz = check_kml_text('orchard.kmz', z.read('doc.kml'))
    n_gj = check_geojson(OUT / 'orchard.geojson')
    check(n_kml == n_kmz == n_gj, f'KML, KMZ and GeoJSON carry the same {n_gj} features')
    check_svg(OUT / 'orchard.svg')
    check_pdf(OUT / 'orchard.pdf', ['Exemplo pomar', 'OpenCADIrrigation'])

    print('report exports')
    check_pdf(OUT / 'orchard_report.pdf', ['Exemplo pomar'])
    check_xlsx(OUT / 'orchard_report.xlsx', 5)
    check_csv(OUT / 'orchard_report.csv')
    check_csv(OUT / 'orchard_schedule.csv')
    txt = (OUT / 'orchard_report.txt').read_text(encoding='utf-8')
    check('OpenCADIrrigation' in txt and len(txt.splitlines()) > 20, f'orchard_report.txt: {len(txt.splitlines())} lines')

    print('project and catalogue files')
    pj = json.loads((OUT / 'orchard.oci.json').read_text(encoding='utf-8'))
    check(pj.get('format') == 'opencadirrigation' and pj.get('entities') and pj.get('results'), f'orchard.oci.json: {len(pj.get("entities", []))} entities with results')
    cat = json.loads((OUT / 'catalog.oci.cat').read_text(encoding='utf-8'))
    check(cat.get('format') == 'opencadirrigation-catalog' and cat.get('catalog', {}).get('pipes'), f'catalog.oci.cat: {cat.get("items")} items')
    check_xlsx(OUT / 'catalog.xlsx', 6)
    check_csv(OUT / 'catalog_pipes.csv')

    print('gallery examples')
    ex = OUT / 'examples'
    for dxf in sorted(ex.glob('*.dxf')):
        if dxf.name.endswith('_libredwg.dxf'):
            continue
        check_dxf(dxf)
        check_dwg(dxf.with_suffix('.dwg'))
        check_kml_text(dxf.stem + '.kml', dxf.with_suffix('.kml').read_bytes())
        check_geojson(dxf.with_suffix('.geojson'))
        check_svg(dxf.with_suffix('.svg'))

    print()
    if failures:
        print(f'{len(failures)} check(s) failed:')
        for f in failures:
            print(' -', f)
        sys.exit(1)
    print('all export checks passed')


if __name__ == '__main__':
    main()
